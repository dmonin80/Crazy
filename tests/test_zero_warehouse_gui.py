import os
from pathlib import Path
import subprocess
import sys


def test_zero_table_row_click_allocate_return_reissue_and_export(tmp_path):
    project=Path(__file__).resolve().parents[1]
    script=r'''
import sys,uuid
from pathlib import Path
from PySide6 import QtCore,QtWidgets as W,QtTest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from aho.client import Client,STOCK_COLUMNS
from aho.core import uid
from aho.server import create_app
from aho.stock_dialog import StockMoveDialog
root=Path(sys.argv[1])
app=W.QApplication([])
api=create_app(root/'data');store=api.state.store;store.setup('admin','stock-test-password')
u=store.user(store.login('Администратор','stock-test-password')['token'])
rows=[{'code':'0001','name':'Стул офисный','unit':'шт','quantity':10,'rows':[2]},{'code':'0002','name':'Монитор','unit':'шт','quantity':2,'rows':[3]}]
store.mutate(u,uid(),'import',{},lambda c:store.import_rows(c,rows,'test-source-hash','test.xlsx'))
def loc(name,parent=None):return store.mutate(u,uid(),'location',{},lambda c:store.add_location(c,{'name':name,'parent':parent}))['id']
from aho.catalogs import create_entry
responsible=store.mutate(u,uid(),'catalog.create:responsibles',{},lambda c:create_entry(c,'responsibles',{'name':'Тестовый ответственный'}))['id']
one=loc('Здание 1');room_one=loc('101',one);two=loc('Здание 2');room_two=loc('201',two)
http=TestClient(api)
class Window(Client):
    def api(self,path,payload=None):
        token=store.login('Администратор','stock-test-password')['token'] if path=='/unused-login' else self.token
        headers={'Authorization':'Bearer '+token}
        result=http.get(path,headers=headers) if payload is None else http.post(path,headers=headers,json=payload)
        if result.is_error:raise RuntimeError(str(result.json().get('detail')))
        return result.json()
window=Window(local_url='http://127.0.0.1:8765',settings_root=root)
window.token=store.login('Администратор','stock-test-password')['token'];window.refresh();window.show();app.processEvents()
assert window.view_tabs.tabText(0)=='Нулевой склад' and window.table.rowCount()==2
assert window.table.columnCount()==len(STOCK_COLUMNS)
assert 'Номенклатура'==window.table.horizontalHeaderItem(0).text()
# Sorting quantities must be numeric: 2 before 10.
window.table.sortItems(4,QtCore.Qt.SortOrder.AscendingOrder)
assert window.table.item(0,4).data(QtCore.Qt.ItemDataRole.DisplayRole)==2
# A row click must resolve the stable ID, even after sorting.
chair=next(r for r in store.stock() if r['code']=='0001')
def table_row(ident):return next(n for n in range(window.table.rowCount()) if window.table.item(n,0).data(QtCore.Qt.ItemDataRole.UserRole)==ident)

def allocate_dialog(dialog):
    assert dialog.stock['id']==chair['id']
    dialog.suffix_table.item(0,0).setText('/0001')
    assert dialog.warehouse.findData('zero')==-1
    dialog.warehouse.setCurrentIndex(dialog.warehouse.findData(one))
    assert dialog.room.findData(room_one)>=0 and dialog.room.findData(room_two)==-1
    dialog.room.setCurrentIndex(dialog.room.findData(room_one))
    assert dialog.payload()[1]['destination']==room_one
    # Switching the main warehouse clears the previously selected child.
    dialog.warehouse.setCurrentIndex(dialog.warehouse.findData(two))
    assert dialog.room.currentData() is None and dialog.room.findData(room_one)==-1
    assert dialog.payload()[1]['destination']==two
    dialog.warehouse.setCurrentIndex(dialog.warehouse.findData(one))
    dialog.room.setCurrentIndex(dialog.room.findData(room_one));dialog.quantity.setValue(2)
    dialog.suffix_table.item(1,0).setText('/0002')
    dialog.responsible.setCurrentIndex(dialog.responsible.findData(responsible));dialog.number.setText('0007')
    dialog.document_date.setDate(QtCore.QDate(2026,10,29))
    dialog.grab().save(str(root/'movement_dialog.png'))
    dialog.submit()
    assert dialog.result is not None,dialog.error.text()
    return W.QDialog.DialogCode.Accepted
original_exec=StockMoveDialog.exec;StockMoveDialog.exec=allocate_dialog
index=table_row(chair['id']);rect=window.table.visualItemRect(window.table.item(index,0))
QtTest.QTest.mouseClick(window.table.viewport(),QtCore.Qt.MouseButton.LeftButton,QtCore.Qt.KeyboardModifier.NoModifier,rect.center());app.processEvents()
assert not store.assets()
next(b for b in window.findChildren(W.QPushButton) if b.text()=='Переместить').click();app.processEvents()
assert len(store.assets())==2 and all(a['item_id']==chair['id'] and a['location']==room_one for a in store.assets())
assert store.stock()[1 if store.stock()[0]['code']=='0002' else 0]['unassigned']==8
assert window.table.item(table_row(chair['id']),7).data(QtCore.Qt.ItemDataRole.DisplayRole)==8
# Cancellation must never create a document or change a balance.
count=len(store.read('SELECT * FROM documents'))
StockMoveDialog.exec=lambda dialog:W.QDialog.DialogCode.Rejected
window.open_stock_action(table_row(chair['id']))
assert len(store.read('SELECT * FROM documents'))==count
# Return an existing asset, then issue the same card from the Zero table.
asset=store.assets()[0];old_id=asset['id'];old_number=asset['number']
payload={'source':room_one,'destination':'zero','assets':[{'id':old_id,'version':asset['version']}],'date':'2026-10-29'}
store.mutate(u,uid(),'move',payload,lambda c:store.movement(c,u['id'],payload))
window.refresh();available=next(r for r in window.stock_rows if r['id']==chair['id']);assert available['available']==9 and available['returned']==1

def returned_dialog(dialog):
    assert len(dialog.returned)==1 and dialog.returned[0]['id']==old_id
    dialog.sources.setCurrentIndex(1)
    dialog.returned_table.item(0,0).setCheckState(QtCore.Qt.CheckState.Checked)
    dialog.warehouse.setCurrentIndex(dialog.warehouse.findData(two))
    dialog.room.setCurrentIndex(dialog.room.findData(room_two))
    path,payload=dialog.payload();assert path=='/move' and payload['assets'][0]['id']==old_id
    assert 'responsible' not in payload
    dialog.submit();assert dialog.result is not None,dialog.error.text()
    return W.QDialog.DialogCode.Accepted
StockMoveDialog.exec=returned_dialog
window.open_stock_action(table_row(chair['id']))
after=next(a for a in store.assets() if a['id']==old_id)
assert after['number']==old_number and after['location']==room_two and after['responsible']=='Тестовый ответственный'
assert len(store.assets())==2
stock=next(r for r in window.stock_rows if r['id']==chair['id']);assert stock['unassigned']==8 and stock['returned']==0 and stock['available']==8
# The same main table is searchable and exports its actual stock columns.
window.search.setText('0001');app.processEvents();assert window.table.rowCount()==1
W.QFileDialog.getSaveFileName=lambda *a,**k:(str(root/'stock.xlsx'),'Excel (*.xlsx)')
window.export();book=load_workbook(root/'stock.xlsx');sheet=book.active
assert sheet.title=='Нулевой склад' and sheet.max_row==2
assert sheet['H1'].alignment.wrap_text and sheet.row_dimensions[1].height==32
assert sheet['B2'].value=='0001' and sheet['H2'].value==8 and sheet['H2'].data_type=='n'
book.close()
window.search.clear();window.grab().save(str(root/'zero_warehouse.png'))
window.view_tabs.setCurrentIndex(1);assert window.table.rowCount()==2
window.allocate();assert window.view_tabs.currentIndex()==0 and window.table.rowCount()==2
# Simulate loss of a successful reply: only query the original request, never send twice.
import httpx
original_api=window.api
sent=[]
def lost_reply(path,payload=None):
    result=original_api(path,payload)
    if path=='/allocate':
        sent.append(payload['request_id'])
        transport=httpx.ReadTimeout('simulated reply loss')
        raise RuntimeError('Результат неопределён') from transport
    return result
window.api=lost_reply
fresh=next(r for r in store.stock() if r['id']==chair['id'])
dialog=StockMoveDialog(window,fresh)
dialog.warehouse.setCurrentIndex(dialog.warehouse.findData(one))
dialog.quantity.setValue(1);dialog.suffix_table.item(0,0).setText('/0003');dialog.submit()
assert dialog.pending is not None and not dialog.submit_button.isEnabled()
dialog.submit();assert len(sent)==1
dialog.check_result();assert dialog.result is not None and len(store.assets())==3
window.api=original_api
StockMoveDialog.exec=original_exec
window.close();print('ZERO_TABLE_WORKFLOW_OK')
'''
    result=subprocess.run([sys.executable,'-c',script,str(tmp_path)],cwd=project,
                          env=dict(os.environ,QT_QPA_PLATFORM='offscreen'),capture_output=True,text=True,timeout=30)
    assert result.returncode==0,result.stdout+'\n'+result.stderr
    assert 'ZERO_TABLE_WORKFLOW_OK' in result.stdout
