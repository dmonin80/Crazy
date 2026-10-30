from concurrent.futures import ThreadPoolExecutor
import json
import os
from pathlib import Path
import subprocess
import sys
import pytest
from aho.core import Store,uid,Conflict
from aho.importer import preview
from aho.quantities import total,subtract,number

@pytest.fixture
def data(tmp_path):
    s=Store(tmp_path);s.setup('admin','decimal-test-password');u=s.user(s.login('Администратор','decimal-test-password')['token'])
    rows=[{'code':'P','name':'Стул','unit':'шт','quantity':4,'rows':[2]}, {'code':'M','name':'Жидкость','unit':'л','quantity':'9.91','rows':[3]}]
    s.mutate(u,uid(),'import',{},lambda c:s.import_rows(c,rows,'hash','test.xlsx'))
    def loc(name):return s.mutate(u,uid(),'location',{},lambda c:s.add_location(c,{'name':name}))['id']
    return s,u,loc('Склад 1'),loc('Склад 2')

def batch(s,u,dest,lines,rid=None):
    p={'destination':dest,'date':'2026-10-29','lines':lines,'floor':'2'}
    return s.mutate(u,rid or uid(),'batch-transfer',p,lambda c:s.batch_movement(c,u['id'],p))

def line(s,code,amount,suffixes):
    return {'kind':'allocate','item_id':next(r for r in s.stock() if r['code']==code)['id'],'quantity':amount,'suffixes':suffixes}

def test_decimal_import_and_duplicate_sum():
    raw='Номенклатура;Код;Ед. изм.;Количество\n  Жидкость  ;0001;л;0,1\n  Жидкость  ;0001;л;0,2\n'.encode()
    r=preview(raw,'test.csv',{'name':0,'code':1,'unit':2,'quantity':3},merge=True)
    assert not r['errors'] and r['rows'][0]['quantity']=='0.3' and r['units']=='0.3'
    assert r['rows'][0]['name']=='  Жидкость  ' and r['rows'][0]['code']=='0001'
    assert r['rows'][0]['tracking_mode']=='quantity'

def test_mixed_atomic_batch_and_idempotency(data):
    s,u,a,b=data;rid=uid()
    lines=[line(s,'P',2,['-1','-2']),line(s,'M','1,25',['-LOT'])]
    first=batch(s,u,a,lines,rid);assert batch(s,u,a,lines,rid)==first
    assert len(s.read('SELECT * FROM documents'))==1 and len(s.assets())==3
    cards=s.assets();assert next(c for c in cards if c['code']=='M')['quantity']=='1.25'
    stock=next(r for r in s.stock() if r['code']=='M');assert stock['unassigned']=='8.66' and stock['allocated']=='1.25'
    # A valid first line must roll back when the next one exceeds availability.
    before=s.stock();documents=len(s.read('SELECT * FROM documents'))
    with pytest.raises(Conflict):batch(s,u,a,[line(s,'P',1,['-3']),line(s,'M',100,['-BAD'])])
    assert s.stock()==before and len(s.read('SELECT * FROM documents'))==documents

def test_decimal_precision_concurrency_and_partial_move(data):
    s,u,a,b=data
    def run(n):
        try:batch(s,u,a,[line(s,'M','6.7',['-R'+str(n)])]);return True
        except Conflict:return False
    with ThreadPoolExecutor(2) as pool:assert sorted(pool.map(run,range(2)))==[False,True]
    assert next(r for r in s.stock() if r['code']=='M')['unassigned']=='3.21'
    card=s.assets()[0];p={'source':a,'destination':b,'date':'2026-10-29','floor':'3','assets':[{'id':card['id'],'version':card['version'],'quantity':'0.7','suffix':'-PART'}]}
    s.mutate(u,uid(),'move',p,lambda c:s.movement(c,u['id'],p))
    old=next(c for c in s.assets() if c['id']==card['id']);new=next(c for c in s.assets() if c['id']!=card['id'])
    assert old['quantity']==6 and old['location']==a and old['number']==card['number']
    assert new['quantity']=='0.7' and new['location']==b and new['floor']=='3'
    assert number(total(c['quantity'] for c in s.assets()))=='6.7'
    sid=s.mutate(u,uid(),'session',{},lambda c:s.start_session(c,{'name':'Проверка'}))['id']
    snaps=[json.loads(r['snapshot']) for r in s.read('SELECT snapshot FROM checks WHERE session_id=?',(sid,))]
    assert number(total(r['quantity'] for r in snaps))=='6.7'

def test_fractional_piece_unit_allowed_without_rounding(data):
    s,u,a,b=data
    batch(s,u,a,[line(s,'P','1.5',['-HALF'])])
    assert s.assets()[0]['quantity']=='1.5'
    assert next(r for r in s.stock() if r['code']=='P')['unassigned']=='2.5'
    # Very small values are exact; they are never rounded to zero.
    batch(s,u,a,[line(s,'M','0.000000000000000000000000000001',['-TINY'])])
    tiny=next(c for c in s.assets() if c['code']=='M')
    assert tiny['quantity']=='0.000000000000000000000000000001'
    assert next(r for r in s.stock() if r['code']=='M')['unassigned']=='9.909999999999999999999999999999'

def test_v2_upgrade_keeps_existing_ids_and_numbers(data):
    s,u,a,b=data;batch(s,u,a,[line(s,'P',1,['-KEEP'])]);card=s.assets()[0]
    with s.transaction() as c:
        c.execute('ALTER TABLE assets DROP COLUMN quantity_decimal');c.execute('ALTER TABLE items DROP COLUMN initial_decimal');c.execute('ALTER TABLE items DROP COLUMN tracking_mode');c.execute("UPDATE meta SET value='2' WHERE key='schema_version'")
    s2=Store(s.root)
    assert s2.assets()[0]['id']==card['id'] and s2.assets()[0]['number']==card['number'] and s2.assets()[0]['quantity']==1
    assert s2.read("SELECT value FROM meta WHERE key='schema_version'")[0]['value']=='3'
    assert list((s.root/'backups').glob('*.sqlite3'))

def test_gui_button_only_batch_and_fractional_transfer(tmp_path):
    project=Path(__file__).resolve().parents[1]
    script=r'''
import sys
from pathlib import Path
from PySide6 import QtCore,QtWidgets as W,QtTest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from aho.core import uid
from aho.server import create_app
from aho.client import Client,COLUMNS
from aho.batch_dialog import BatchStockDialog
from aho.stock_dialog import StockMoveDialog
from aho.movement_dialog import MoveAssetsDialog
app=W.QApplication([]);root=Path(sys.argv[1]);server=create_app(root/'data');s=server.state.store
s.setup('admin','batch-gui-password');token=s.login('Администратор','batch-gui-password')['token'];u=s.user(token)
s.mutate(u,uid(),'import',{},lambda c:s.import_rows(c,[{'code':'P','name':'Стул','unit':'шт','quantity':4,'rows':[2]},{'code':'M','name':'Жидкость','unit':'л','quantity':'9.91','rows':[3]}],'hash','test.xlsx'))
def loc(name,parent=None):return s.mutate(u,uid(),'location',{},lambda c:s.add_location(c,{'name':name,'parent':parent}))['id']
warehouse=loc('Здание');room=loc('201',warehouse);other=loc('Архив',warehouse)
http=TestClient(server)
class Window(Client):
    def api(self,path,payload=None):
        h={'Authorization':'Bearer '+self.token}
        r=http.get(path,headers=h) if payload is None else http.post(path,headers=h,json=payload)
        if r.is_error:raise RuntimeError(str(r.json().get('detail')))
        return r.json()
w=Window(local_url='http://127.0.0.1:8765',settings_root=root);w.token=token;w.refresh();w.show();app.processEvents()
opened=[]
original_batch=BatchStockDialog.exec
original_single=StockMoveDialog.exec
StockMoveDialog.exec=lambda d:(_ for _ in ()).throw(AssertionError('Row click opened movement'))
for index in range(2):
    rect=w.table.visualItemRect(w.table.item(index,0))
    QtTest.QTest.mouseClick(w.table.viewport(),QtCore.Qt.MouseButton.LeftButton,QtCore.Qt.KeyboardModifier.ControlModifier if index else QtCore.Qt.KeyboardModifier.NoModifier,rect.center())
app.processEvents();assert len(w.table.selectionModel().selectedRows())==2 and not s.assets()
w.grab().save(str(root/'selected_rows.png'))
def drive(d):
    opened.append(True);assert len(d.editors)==2
    d.warehouse.setCurrentIndex(d.warehouse.findData(warehouse));d.room.setCurrentIndex(d.room.findData(room));d.floor.setText('2')
    for e in d.editors:
        if e.stock['code']=='P':
            e.quantity.setValue(2);e.suffix_table.item(0,0).setText('-A');e.suffix_table.item(1,0).setText('-B')
        else:
            e.quantity.setValue('0,35');e.suffix_table.item(0,0).setText('-M')
    d.show();app.processEvents();d.grab().save(str(root/'batch_dialog.png'))
    d.submit();assert d.result is not None,d.error.text()
    return W.QDialog.DialogCode.Accepted
BatchStockDialog.exec=drive
next(b for b in w.findChildren(W.QPushButton) if b.text()=='Переместить').click()
assert opened==[True] and len(s.read('SELECT * FROM documents'))==1 and len(s.assets())==3
m=next(a for a in s.assets() if a['code']=='M');assert m['quantity']=='0.35'
assert next(r for r in s.stock() if r['code']=='M')['unassigned']=='9.56'
# The material card can be moved in part; the original ID identifies the remainder.
d=MoveAssetsDialog(w,[m]);d.destination=other;d.amounts.item(0,1).setText('0.1');d.amounts.item(0,2).setText('-PART');d.submit()
assert d.result()==W.QDialog.DialogCode.Accepted,d.error.text()
assert next(a for a in s.assets() if a['id']==m['id'])['quantity']=='0.25'
w.refresh();w.view_tabs.setCurrentIndex(1)
quantity_column=next(i for i,(key,_) in enumerate(COLUMNS) if key=='quantity')
assert {w.table.item(n,quantity_column).text() for n in range(w.table.rowCount())}=={'1','0.25','0.1'}
w.view_tabs.setCurrentIndex(0)
W.QFileDialog.getSaveFileName=lambda *a,**k:(str(root/'stock.xlsx'),'Excel (*.xlsx)')
w.export();book=load_workbook(root/'stock.xlsx');ws=book.active
material=next(n for n in range(2,ws.max_row+1) if ws.cell(n,2).value=='M')
assert ws.cell(material,5).value==9.56 and ws.cell(material,5).data_type=='n'
assert Client.export_value('quantity','0.123456789012345678901')=='0.123456789012345678901'
book.close();BatchStockDialog.exec=original_batch;StockMoveDialog.exec=original_single;w.close()
print('BUTTON_BATCH_DECIMAL_OK')
'''
    result=subprocess.run([sys.executable,'-c',script,str(tmp_path)],cwd=project,env=dict(os.environ,QT_QPA_PLATFORM='offscreen'),capture_output=True,text=True,timeout=30)
    assert result.returncode==0,result.stdout+'\n'+result.stderr
    assert 'BUTTON_BATCH_DECIMAL_OK' in result.stdout


def test_computed_totals_are_exact_and_not_limited_as_one_input():
    from decimal import Decimal
    assert number(total(['100000000000000000000000000000']*20))==2000000000000000000000000000000
    assert number(total(['0.1','0.2']))=='0.3'
    assert number(subtract('1','0.000000000000000000000000000001'))=='0.999999999999999999999999999999'


def test_large_decimal_import_does_not_overflow_sqlite(tmp_path):
    s=Store(tmp_path);s.setup('admin','large-quantity-password');u=s.user(s.login('Администратор','large-quantity-password')['token'])
    rows=[{'code':'L','name':'Материал','unit':'л','quantity':'100000000000000000000','rows':[2]}]
    s.mutate(u,uid(),'import',{},lambda c:s.import_rows(c,rows,'large-hash','test.xlsx'))
    assert s.stock()[0]['initial']==100000000000000000000 and s.stock()[0]['unassigned']==100000000000000000000
