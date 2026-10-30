from io import BytesIO
import base64
import json
import os
from pathlib import Path
import subprocess
import sys
import uuid

from fastapi.testclient import TestClient
from openpyxl import Workbook
import pytest

from aho.importer import inspect_source,preview
from aho.server import create_app

MAP={'name':8,'code':7,'unit':10,'quantity':12}


def fixture_xlsx():
    wb=Workbook();ws=wb.active;ws.title='Выгрузка 1С'
    for column,value in {1:'Счёт',2:'Подразделение',5:'Склад',7:'МОЛ',8:'Код',9:'Номенклатура',10:'Цена',11:'Единица измерения',12:'Сумма',13:'Количество'}.items():
        ws.cell(1,column,value)
    def row(number,code,name,unit,count):
        for col,value in ((8,code),(9,name),(11,unit),(13,count)):
            if value is not None:ws.cell(number,col,value)
    row(2,'00001','  Стул офисный  ','шт',2)
    row(3,'00001','  Стул офисный  ','шт',3)
    row(8,'TEST-FUEL','Тестовое топливо','л',9.91)
    row(197,'TEST-METAL','Тестовый материал','т',None)
    ws.merge_cells('E2:F2')
    ws['J2']='=1+1'  # No cached result in an irrelevant column must not block import.
    out=BytesIO();wb.save(out);return out.getvalue()


def test_detect_real_1c_column_layout():
    i=inspect_source(fixture_xlsx(),'1c.xlsx')
    assert i['header']==1 and i['mapping']==MAP
    assert [i['columns'][MAP[k]]['letter'] for k in ('name','code','unit','quantity')]==['I','H','K','M']
    r=preview(fixture_xlsx(),'1c.xlsx',MAP,merge=True)
    assert [(e['row'],e['cell'],e['field']) for e in r['errors']]==[(197,'M197','Количество')]
    assert r['rows'][0]['name']=='  Стул офисный  '
    assert r['rows'][0]['code']=='00001'


def test_explicit_exclusions_and_merge_only():
    raw=fixture_xlsx()
    result=preview(raw,'1c.xlsx',MAP,merge=True,excluded_rows=[8,197])
    assert not result['errors'] and result['units']==5
    assert result['rows'][0]['rows']==[2,3]
    assert [r['row'] for r in result['excluded']]==[8,197]
    assert preview(raw,'1c.xlsx',MAP,excluded_rows=[8,197])['errors']
    assert len(preview(raw,'1c.xlsx',MAP,merge=True)['errors'])==1
    for excluded in ([1],[999],[8,8],[-1],['8']):
        with pytest.raises(ValueError):preview(raw,'1c.xlsx',MAP,excluded_rows=excluded)


def test_only_clear_service_rows_skipped():
    data='Номенклатура;Код;Ед. изм.;Кол-во\nСтул;00001;шт;2\nНоменклатура;Код;Ед. изм.;Кол-во\nИтого;;;2\nИтого;ITEM-2;шт;1\nПозиция без количества;ITEM-3;шт;\n'.encode()
    i=inspect_source(data,'1c.csv');assert len(i['mapping'])==4
    r=preview(data,'1c.csv',i['mapping'])
    assert [s['row'] for s in r['skipped']]==[3,4]
    assert [s['code'] for s in r['rows']]==['00001','ITEM-2']
    assert r['errors'][0]['cell']=='D6'


def test_ragged_csv_and_invalid_mapping():
    data='Номенклатура;Код;Ед. изм.;Количество\nСтул;A;шт;2\nСтол;B\n'.encode()
    mapping={'name':0,'code':1,'unit':2,'quantity':3}
    r=preview(data,'1c.csv',mapping)
    assert {e['cell'] for e in r['errors']}=={'C3','D3'}
    with pytest.raises(ValueError):preview(data,'1c.csv',dict(mapping,code=999))


def test_inspect_nonfirst_header_sheet_and_encoding():
    wb=Workbook();wb.active.title='Обложка';wb.active['A1']='Только обложка'
    ws=wb.create_sheet('Данные');ws.append(['Отчёт']);ws.append(['наименование','код 1С','Ед. изм.','Кол-во']);ws.append(['Стул','001','шт',2]);out=BytesIO();wb.save(out)
    i=inspect_source(out.getvalue(),'a.xlsx')
    assert i['sheet']=='Данные' and i['header']==2 and len(i['mapping'])==4
    data='Номенклатура;Код;Ед. изм.;Количество\nСтул;001;шт;2'.encode('cp1251')
    assert inspect_source(data,'a.csv',encoding='cp1251')['mapping']['quantity']==3
    with pytest.raises(ValueError):inspect_source(data,'a.csv',encoding='utf-8-sig')


def test_import_exclusions_are_audited_and_source_preserved(tmp_path):
    app=create_app(tmp_path);s=app.state.store;s.setup('admin','import-test-password')
    with TestClient(app) as client:
        token=client.post('/login',json={'name':'Администратор','password':'import-test-password'}).json()['token'];headers={'Authorization':'Bearer '+token}
        raw=fixture_xlsx();p={'data':base64.b64encode(raw).decode(),'filename':'1c.xlsx','mapping':MAP,'header':1,'merge':True,'excluded_rows':[8,197],'request_id':str(uuid.uuid4())}
        inspect=client.post('/imports/inspect',headers=headers,json=p)
        assert inspect.status_code==200 and inspect.json()['mapping']==MAP
        r=client.post('/imports/commit',headers=headers,json=p)
        assert r.status_code==200 and r.json()=={'positions':1,'units':5}
        stock=client.get('/stock',headers=headers).json();assert stock[0]['initial']==5
        audit=s.read("SELECT data FROM audit WHERE action='import'")
        assert json.loads(audit[0]['data'])['excluded_rows']==[8,197]
        assert next((tmp_path/'imports').glob('*.xlsx')).read_bytes()==raw


def test_wizard_gui_selection_validation_and_commit(tmp_path):
    project=Path(__file__).resolve().parents[1]
    source=tmp_path/'source.xlsx';source.write_bytes(fixture_xlsx())
    script=r'''
import sys,uuid
from pathlib import Path
from PySide6 import QtCore,QtWidgets as W
from fastapi.testclient import TestClient
from aho.server import create_app
from aho.import_dialog import ImportDialog
app=W.QApplication([])
root=Path(sys.argv[1]);source=Path(sys.argv[2])
server=create_app(root/'data');server.state.store.setup('admin','wizard-test-password')
http=TestClient(server)
token=http.post('/login',json={'name':'Администратор','password':'wizard-test-password'}).json()['token']
class Parent(W.QWidget):
    def api(self,path,payload=None):
        h={'Authorization':'Bearer '+token}
        r=http.get(path,headers=h) if payload is None else http.post(path,headers=h,json=payload)
        assert r.status_code==200,r.text
        return r.json()
    def command(self,path,p):return self.api(path,dict(p,request_id=str(uuid.uuid4())))
    def refresh(self):self.refreshed=True
parent=Parent();parent.refreshed=False
dialog=ImportDialog(parent,source);dialog.show();app.processEvents()
assert {k:c.currentData() for k,c in dialog.mapping.items()}=={'name':8,'code':7,'unit':10,'quantity':12}
assert not dialog.import_button.isEnabled()
dialog.merge.setChecked(True);dialog.validate()
assert dialog.issues.rowCount()==1
for row in range(1):dialog.issues.item(row,0).setCheckState(QtCore.Qt.CheckState.Checked)
W.QMessageBox.question=lambda *a,**k:W.QMessageBox.StandardButton.Yes
dialog.exclude_checked();assert dialog.excluded=={197}
assert dialog.import_button.isEnabled()
assert dialog.result['units']=='14.91'
dialog.grab().save(str(root/'import_preview.png'))
dialog.commit()
assert parent.refreshed
assert server.state.store.stock()[0]['initial']==5
assert dialog.result['rows'][0]['name']=='  Стул офисный  '
print('WIZARD_COMMIT_OK')
'''
    result=subprocess.run([sys.executable,'-c',script,str(tmp_path),str(source)],cwd=project,env=dict(os.environ,QT_QPA_PLATFORM='offscreen'),capture_output=True,text=True,timeout=25)
    assert result.returncode==0,result.stdout+'\n'+result.stderr
    assert 'WIZARD_COMMIT_OK' in result.stdout
