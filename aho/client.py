import json
import sys
import uuid
import hashlib
from datetime import datetime
from urllib.parse import urlsplit
from pathlib import Path
import httpx
from PySide6 import QtCore, QtWidgets as W
from .quantities import quantity, number as qnumber, total as qtotal

STOCK_COLUMNS = [
    ("name", "Номенклатура"),
    ("code", "Код 1С"),
    ("unit", "Ед. изм."),
    ("initial", "Начальное\nколичество"),
    ("unassigned", "Нераспределено"),
    ("returned", "Возвращено"),
    ("allocated", "На других\nскладах"),
    ("available", "Доступно на\nНулевом складе"),
]

COLUMNS = [
    ("name", "Наименование"),
    ("number", "Инвентарный номер"),
    ("quantity", "Общее количество"),
    ("object", "Объект"),
    ("floor", "Этаж"),
    ("room", "Кабинет"),
    ("responsible", "Ответственное лицо"),
    ("location_name", "Наименование склада"),
    ("note", "Примечание"),
    ("supplier", "Поставщик"),
]


class APIError(RuntimeError):
    def __init__(self, message, status):
        super().__init__(message)
        self.status = status


class PendingOperation(RuntimeError):
    pass


class Client(W.QMainWindow):
    def __init__(self, local_url=None, settings_root=None, server_url=None):
        super().__init__()
        self.setWindowTitle("АХО — система инвентаризации · 0.3.1")
        self.resize(1400, 800)
        root = (
            Path(settings_root)
            if settings_root
            else (
                Path(sys.executable).parent
                if getattr(sys, "frozen", False)
                else Path(__file__).resolve().parent.parent
            )
        )
        self.preferences = QtCore.QSettings("AHO", "Inventory")
        self._busy = False
        self._http = None
        self._http_key = None
        self._close_requested = False
        self._last_sync = None
        self._online = False
        self._header_view = None
        self.server_url = server_url
        self.local_url = local_url
        self.settings_path = root / "client_settings.json"
        self.settings = (
            {"server": server_url}
            if server_url
            else (
                json.loads(self.settings_path.read_text("utf-8"))
                if self.settings_path.exists()
                else {"server": "http://127.0.0.1:8765"}
            )
        )
        if local_url:
            self.settings["server"] = local_url
        self.column_filters = {0: {}, 1: {}}
        self.current_user = {}
        self.buttons = {}
        self.token = ""
        self.rows = []
        self.stock_rows = []
        self.locations = []
        self.sid = None
        self.visible = []
        self._stock_dialog_active = False
        self._displayed_view = None
        box = W.QWidget()
        self.setCentralWidget(box)
        layout = W.QVBoxLayout(box)
        title = W.QLabel("Система инвентаризации имущества")
        title.setStyleSheet("font-size:24px;font-weight:600;padding:16px;color:#17436b")
        layout.addWidget(title)
        self.user_label = W.QLabel("Пользователь: вход не выполнен")
        self.user_label.setStyleSheet(
            "font-size:14px;font-weight:600;color:#17436b;padding:8px 18px"
        )
        self.menuBar().setCornerWidget(self.user_label, QtCore.Qt.Corner.TopRightCorner)
        from .view_switcher import ViewSwitcher

        self.view_tabs = ViewSwitcher()
        layout.addWidget(self.view_tabs)
        bar = W.QHBoxLayout()
        layout.addLayout(bar)
        for name, fn in [
            ("Вход" if local_url or server_url else "Подключение", self.login),
            ("Пользователи", self.open_users),
            ("Обновить", self.refresh),
            ("Импорт 1С", self.import_file),
            ("Справочники", self.open_directories),
            ("Переместить", self.move),
            ("Вернуть", lambda: self.move(True)),
            ("Инвентаризация", self.inventory),
            ("Этикетки PDF", self.labels),
            ("Экспорт XLSX", self.export),
        ]:
            button = W.QPushButton(name)
            button.clicked.connect(lambda checked=False, f=fn: self.safe(f))
            bar.addWidget(button)
            self.buttons[name] = button
        self.search = W.QLineEdit()
        self.search.setPlaceholderText(
            "Поиск по наименованию, коду и другим значениям таблицы"
        )
        self.search.textChanged.connect(self.populate)
        layout.addWidget(self.search)
        from .table_model import InventoryTable

        self.table = InventoryTable()
        self.table.replace([], COLUMNS)
        self.table.setSelectionBehavior(
            W.QAbstractItemView.SelectionBehavior.SelectRows
        )
        self.table.setSelectionMode(W.QAbstractItemView.SelectionMode.ExtendedSelection)
        self.table.setEditTriggers(W.QAbstractItemView.EditTrigger.NoEditTriggers)
        self.table.setSortingEnabled(True)
        self.table.horizontalHeader().setSectionsMovable(True)
        self.table.setWordWrap(False)
        self.table.verticalHeader().setDefaultSectionSize(32)
        from .column_filter import FilterHeader

        header = FilterHeader(self.table)
        self.table.setHorizontalHeader(header)
        header.setSectionsMovable(True)
        header.filterRequested.connect(
            lambda col: self.safe(lambda: self.open_column_filter(col))
        )
        header.setToolTip(
            "Нажмите ▼ или правой кнопкой по заголовку для фильтра. Нажатие на название — сортировка."
        )
        self.table.doubleClicked.connect(
            lambda _: self.safe(self.card)
            if self.view_tabs.currentIndex() == 1
            else None
        )
        layout.addWidget(self.table)
        self.table.itemSelectionChanged.connect(self.update_selection)
        self.filter_bar = W.QHBoxLayout()
        layout.addLayout(self.filter_bar)
        self.filter_label = W.QLabel()
        self.filter_bar.addWidget(self.filter_label, 1)
        self.filter_chips = W.QHBoxLayout()
        self.filter_bar.addLayout(self.filter_chips)
        preset = W.QPushButton("Сохранённые фильтры…")
        preset.clicked.connect(self.filter_presets)
        self.filter_bar.addWidget(preset)
        reset = W.QPushButton("Сбросить фильтры")
        reset.clicked.connect(self.clear_filters)
        self.filter_bar.addWidget(reset)
        columns_button = W.QPushButton("Колонки…")
        columns_button.clicked.connect(self.configure_columns)
        self.filter_bar.addWidget(columns_button)
        self.view_tabs.currentChanged.connect(lambda _: self.populate())
        scanbar = W.QHBoxLayout()
        layout.addLayout(scanbar)
        self.scan = W.QLineEdit()
        self.scan.setPlaceholderText("QR-код или инвентарный номер → Enter")
        self.scan.returnPressed.connect(lambda: self.safe(self.scan_code))
        scanbar.addWidget(self.scan)
        self.quick = W.QCheckBox("Быстрая отметка в выбранной инвентаризации")
        self.quick.hide()
        self.session_label = W.QLabel(
            "Сканирование и отметки — в разделе «Инвентаризации»"
        )
        scanbar.addWidget(self.session_label)
        self.connection_label = W.QLabel("Нет подключения")
        self.statusBar().addPermanentWidget(self.connection_label)
        self.security_banner = W.QLabel()
        self.security_banner.setWordWrap(True)
        self.security_banner.setStyleSheet(
            "padding:8px;background:#fff0bc;color:#674b00"
        )
        layout.insertWidget(1, self.security_banner)
        self.update_connection()
        navigation = self.menuBar().addMenu("Разделы")
        navigation.addAction("Остатки", lambda: self.view_tabs.setCurrentIndex(0))
        navigation.addAction("Имущество", lambda: self.view_tabs.setCurrentIndex(1))
        navigation.addAction("Инвентаризации", lambda: self.safe(self.inventory))
        navigation.addAction("Документы", lambda: self.safe(self.open_documents))
        navigation.addAction("Справочники", lambda: self.safe(self.open_directories))
        admin_menu = self.menuBar().addMenu("Управление")
        admin_menu.addAction("Пользователи", lambda: self.safe(self.open_users))
        for caption in ("Пользователи", "Справочники"):
            self.buttons[caption].hide()
        system = self.menuBar().addMenu("Сервис")
        system.addAction(
            "Незавершённые операции", lambda: self.safe(self.pending_operations)
        )
        system.addAction(
            "Интервал обновления…", lambda: self.safe(self.refresh_settings)
        )
        system.addAction("Резервные копии", lambda: self.safe(self.create_backup))
        system.addAction("Журнал действий", lambda: self.safe(self.show_audit))
        self.statusBar().showMessage("Подключитесь к серверу")
        self.setStyleSheet(
            "QPushButton {padding:8px;border:1px solid #b6c7d6;border-radius:5px;background:#edf4fa} QTableView {alternate-background-color:#edf4fa;gridline-color:#dbe3eb;border:1px solid #dbe3eb} QLineEdit {padding:6px} QHeaderView::section {padding:7px;background:#f0f4f8;border:0;border-right:1px solid #dbe3eb;font-weight:600}"
        )
        self.table.setAlternatingRowColors(True)
        self.populate()
        for label, button in self.buttons.items():
            if label not in ("Вход", "Подключение"):
                button.setEnabled(False)

    def safe(self, fn):
        if self._busy:
            return
        try:
            fn()
        except Exception as e:
            W.QMessageBox.warning(self, "Операция не выполнена", str(e))

    def api(self, path, payload=None):
        if self._busy:
            raise RuntimeError("Дождитесь завершения текущей операции")
        operation_key = None
        if payload is not None and "request_id" in payload:
            operation_key = self._pending_key(
                path, {k: v for k, v in payload.items() if k != "request_id"}
            )
            saved = self.preferences.value(operation_key, "")
            if saved:
                payload["request_id"] = json.loads(saved)["id"]
            else:
                record = {
                    "id": payload["request_id"],
                    "path": path,
                    "server": self.settings["server"],
                    "user": self.current_user.get("id"),
                    "at": datetime.now().isoformat(timespec="seconds"),
                }
                self.preferences.setValue(operation_key, json.dumps(record))
                self.preferences.sync()
        self._busy = True
        from .tasks import wait_for_io

        try:
            result = wait_for_io(self, lambda: self._request(path, payload))
            self._online = True
            if operation_key:
                self.preferences.remove(operation_key)
            return result
        except APIError as exc:
            if operation_key and 400 <= exc.status < 500:
                self.preferences.remove(operation_key)
            raise
        except httpx.TransportError as exc:
            self._online = False
            raise PendingOperation(
                "Связь прервана. Данные могут быть неактуальны. Результат операции будет проверен при повторе."
            ) from exc
        finally:
            self._busy = False
            self.update_connection()

    def _request(self, path, payload=None):
        key = (
            self.settings["server"],
            self.settings.get("ca_file"),
            bool(self.local_url),
        )
        if self._http is None or self._http_key != key:
            if self._http:
                self._http.close()
            self._http = httpx.Client(
                timeout=httpx.Timeout(30, connect=5),
                verify=True
                if self.local_url
                else (self.settings.get("ca_file") or True),
                trust_env=False,
            )
            self._http_key = key
        headers = {"Authorization": "Bearer " + self.token} if self.token else {}
        response = self._http.request(
            "GET" if payload is None else "POST",
            self.settings["server"].rstrip("/") + path,
            json=payload,
            headers=headers,
        )
        if response.is_error:
            try:
                message = response.json().get("detail", response.text)
            except ValueError:
                message = response.text
            if response.status_code == 403 and message == "Требуется вход":
                message = "Сеанс завершён. Войдите в программу снова."
            raise APIError(str(message), response.status_code)
        return (
            response.content
            if "application/pdf" in response.headers.get("content-type", "")
            else response.json()
        )

    def _pending_key(self, path, payload):
        raw = json.dumps(
            [self.settings["server"], self.current_user.get("id"), path, payload],
            sort_keys=True,
            ensure_ascii=False,
        )
        return "pending/" + hashlib.sha256(raw.encode()).hexdigest()

    def command(self, path, payload):
        key = self._pending_key(path, payload)
        saved = self.preferences.value(key, "")
        record = (
            json.loads(saved)
            if saved
            else {
                "id": str(uuid.uuid4()),
                "path": path,
                "server": self.settings["server"],
                "user": self.current_user.get("id"),
                "at": datetime.now().isoformat(timespec="seconds"),
            }
        )
        self.preferences.setValue(key, json.dumps(record))
        self.preferences.sync()

        def stored_result():
            try:
                return True, self.api("/requests/" + record["id"])
            except APIError as exc:
                if exc.status == 404:
                    return False, None
                raise

        if saved:
            found, result = stored_result()
            if found:
                self.preferences.remove(key)
                return result
        try:
            result = self.api(path, dict(payload, request_id=record["id"]))
        except APIError as exc:
            if 400 <= exc.status < 500:
                self.preferences.remove(key)
            raise
        except PendingOperation:
            found, result = stored_result()
            if not found:
                raise PendingOperation(
                    "Ответ не получен. Операция сохранена для проверки. Повторите то же действие: программа использует прежний номер запроса и не проведёт его дважды."
                )
        self.preferences.remove(key)
        return result

    def update_connection(self):
        if not hasattr(self, "connection_label"):
            return
        stamp = (
            self._last_sync.strftime("%H:%M:%S") if self._last_sync else "ещё не было"
        )
        self.connection_label.setText(
            ("● На связи" if self._online else "● Нет связи") + " · Обновлено: " + stamp
        )
        self.connection_label.setStyleSheet(
            "color:" + ("#17623b" if self._online else "#ae2020")
        )
        address = urlsplit(self.settings.get("server", ""))
        insecure = address.scheme == "http" and address.hostname not in (
            "127.0.0.1",
            "localhost",
            "::1",
        )
        self.security_banner.setText(
            "Незащищённое сетевое подключение HTTP: пароль и данные передаются без шифрования. Настройте HTTPS на сервере."
        )
        self.security_banner.setVisible(insecure)

    def pending_operations(self):
        rows = []
        for key in self.preferences.allKeys():
            if not key.startswith("pending/"):
                continue
            record = json.loads(self.preferences.value(key))
            if record["server"] != self.settings["server"] or record[
                "user"
            ] != self.current_user.get("id"):
                continue
            try:
                self.api("/requests/" + record["id"])
                rows.append(record["at"] + " · " + record["path"] + " · Проведено")
                self.preferences.remove(key)
            except APIError as exc:
                if exc.status != 404:
                    raise
                rows.append(
                    record["at"]
                    + " · "
                    + record["path"]
                    + " · Не подтверждено; повторите исходное действие"
                )
        W.QMessageBox.information(
            self,
            "Результаты операций",
            "\n".join(rows) or "Незавершённых операций нет.",
        )

    def closeEvent(self, event):
        if self._busy or (
            self._close_requested
            and W.QApplication.activeModalWidget() not in (None, self)
        ):
            self._close_requested = True
            QtCore.QTimer.singleShot(100, self._complete_close)
            event.ignore()
            return
        self._close_requested = False
        self.save_table_state()
        if self._http:
            self._http.close()
            self._http = None
        super().closeEvent(event)

    def _complete_close(self):
        if self._close_requested:
            self.close()

    def text(self, title, label, default=""):
        value, ok = W.QInputDialog.getText(self, title, label, text=default)
        if not ok:
            raise ValueError("Операция отменена")
        return value

    def choose_location(self, allow_zero=True):
        from .reference_widgets import LocationDialog

        self.locations = self.api("/locations")
        dialog = LocationDialog(self, allow_zero)
        if dialog.exec() != W.QDialog.DialogCode.Accepted:
            raise ValueError("Операция отменена")
        return dialog.ident

    def open_directories(self):
        from .directory_dialog import DirectoriesDialog

        DirectoriesDialog(self).exec()
        self.refresh()

    def login_with_credentials(self, name, password):
        result = self.api("/login", {"name": name, "password": password})
        self.token = result["token"]
        self.current_user = self.api("/me")
        self.apply_permissions()
        if not self.server_url:
            self.settings_path.write_text(
                json.dumps(self.settings, ensure_ascii=False, indent=2), "utf-8"
            )
        self.refresh()
        self.statusBar().showMessage(
            ("Локальный сервер · " if self.local_url else "Подключено · ")
            + name
            + " · "
            + result["role"]
        )

    def login(self):
        if self.local_url or self.server_url:
            server = self.local_url or self.server_url
        else:
            server = self.text("Подключение", "Адрес сервера", self.settings["server"])
        self.settings["server"] = server
        from .user_dialog import LoginDialog

        dialog = LoginDialog(self)
        if dialog.exec() != W.QDialog.DialogCode.Accepted and not self.token:
            self.close()

    def apply_permissions(self):
        u = self.current_user
        self.user_label.setText("Пользователь: " + u.get("name", "вход не выполнен"))
        mapping = {
            "Импорт 1С": "import",
            "Справочники": "catalogs",
            "Переместить": "move",
            "Вернуть": "return",
            "Инвентаризация": "inventory",
            "Этикетки PDF": "labels",
            "Экспорт XLSX": "export",
        }
        for label, permission in mapping.items():
            self.buttons[label].setEnabled(
                u.get("is_admin") or permission in u.get("permissions", [])
            )
        self.buttons["Переместить"].setEnabled(
            bool(
                u.get("is_admin")
                or any(p in u.get("permissions", []) for p in ("allocate", "move"))
            )
        )
        self.buttons["Пользователи"].setEnabled(bool(u.get("is_admin")))
        self.buttons["Обновить"].setEnabled(bool(u))
        self.update_selection()

    def open_users(self):
        from .user_dialog import UsersDialog

        UsersDialog(self).exec()

    def refresh(self):
        if type(self).api is Client.api:
            overview = self.api("/overview")
            self.current_user = overview["user"]
            rows = overview["assets"]
            stock = overview["stock"]
        else:
            self.current_user = self.api("/me")
            rows = self.api("/assets")
            stock = self.api("/stock")
        self.apply_permissions()
        self.rows = rows
        self.stock_rows = [
            dict(r, available=qnumber(qtotal([r["unassigned"], r["returned"]])))
            for r in stock
        ]
        self._last_sync = datetime.now()
        self._online = True
        self.update_connection()
        self.populate()

    def populate(self):
        zero = self.view_tabs.currentIndex() == 0
        columns = STOCK_COLUMNS if zero else COLUMNS
        source = self.stock_rows if zero else self.rows
        query = self.search.text().strip().casefold()
        from .column_filter import matches

        filters = self.column_filters[self.view_tabs.currentIndex()]
        self.visible = [
            r
            for r in source
            if matches(r, filters)
            and query in " ".join(str(r.get(key, "")) for key, _ in columns).casefold()
        ]
        scroll = self.table.verticalScrollBar().value()
        selected_ids = {
            self.table.item(i.row(), 0).data(QtCore.Qt.ItemDataRole.UserRole)
            for i in self.table.selectionModel().selectedRows()
            if self.table.item(i.row(), 0)
        }
        if self._header_view is not None and self._header_view != zero:
            self.save_table_state()
        self.table.setSortingEnabled(False)
        self.table.blockSignals(True)
        self.table.replace(self.visible, columns)
        self.table.horizontalHeader().active = {
            i for i, (key, _) in enumerate(columns) if key in filters
        }

        if self._displayed_view != zero:
            widths = (
                [350, 145, 85, 135, 140, 110, 135, 165]
                if zero
                else [250, 165, 140, 185, 130, 110, 185, 230, 180, 190]
            )
            for col, width in enumerate(widths):
                self.table.setColumnWidth(col, width)
            self.table.horizontalHeader().setMinimumHeight(44 if zero else 28)
            self.table.horizontalHeader().setSortIndicator(
                0, QtCore.Qt.SortOrder.AscendingOrder
            )
            selected_ids = set()
            scroll = 0
            self._displayed_view = zero
        if self._header_view != zero:
            state = self.preferences.value("table/" + ("stock" if zero else "assets"))
            if state:
                self.table.horizontalHeader().restoreState(state)
            self._header_view = zero
        self.table.setSortingEnabled(True)
        self.table.clearSelection()
        for row in range(self.table.rowCount()):
            if (
                self.table.item(row, 0).data(QtCore.Qt.ItemDataRole.UserRole)
                in selected_ids
            ):
                self.table.selectionModel().select(
                    self.table.model().index(row, 0),
                    QtCore.QItemSelectionModel.SelectionFlag.Select
                    | QtCore.QItemSelectionModel.SelectionFlag.Rows,
                )
        self.table.blockSignals(False)
        self.update_selection()
        while self.filter_chips.count():
            item = self.filter_chips.takeAt(0)
            if item.widget():
                item.widget().deleteLater()
        for key, label in columns:
            if key in filters:
                chip = W.QPushButton(label.replace("\n", " ") + " ×")
                chip.setToolTip(", ".join(sorted(filters[key])))
                chip.clicked.connect(lambda checked=False, k=key: self.remove_filter(k))
                self.filter_chips.addWidget(chip)
        self.filter_label.setText(
            "Показано "
            + str(len(self.visible))
            + " из "
            + str(len(source))
            + (
                " · Фильтры: "
                + ", ".join(
                    label.replace("\n", " ") for key, label in columns if key in filters
                )
                if filters
                else ""
            )
        )
        self.table.horizontalHeader().viewport().update()
        self.table.verticalScrollBar().setValue(scroll)
        self.statusBar().showMessage(
            ("Нулевой склад · позиций: " if zero else "Найдено карточек: ")
            + str(len(self.visible))
            + (
                " · Выделение: Ctrl / Shift · Перенос — кнопка «Переместить»"
                if zero
                else ""
            )
        )

    def open_column_filter(self, col):
        from .column_filter import ColumnFilterDialog, value_text, matches

        view = self.view_tabs.currentIndex()
        columns = STOCK_COLUMNS if view == 0 else COLUMNS
        if not 0 <= col < len(columns):
            return
        key, label = columns[col]
        filters = self.column_filters[view]
        source = self.stock_rows if view == 0 else self.rows
        query = self.search.text().strip().casefold()
        values = {
            value_text(r, key)
            for r in source
            if matches(r, filters, exclude=key)
            and query in " ".join(str(r.get(key, "")) for key, _ in columns).casefold()
        }
        selected = filters.get(key)
        # Keep selected values even when other filters currently hide them.
        dialog = ColumnFilterDialog(self, label, values | (selected or set()), selected)
        if dialog.exec() != W.QDialog.DialogCode.Accepted:
            return
        if dialog.sort_order is not None:
            self.table.sortItems(col, dialog.sort_order)
            return
        if dialog.result_values is None:
            filters.pop(key, None)
        else:
            filters[key] = dialog.result_values
        self.populate()

    def selected(self):
        if self.view_tabs.currentIndex() == 0:
            return []
        ids = {
            self.table.item(i.row(), 0).data(QtCore.Qt.ItemDataRole.UserRole)
            for i in self.table.selectionModel().selectedRows()
        }
        return [r for r in self.rows if r["id"] in ids]

    def open_stock_action(self, row=None):
        if self._stock_dialog_active:
            return
        indices = [i.row() for i in self.table.selectionModel().selectedRows()]
        if row is not None:
            indices = [row]
        if not indices:
            raise ValueError("Выделите одно или несколько наименований для перемещения")
        ids = {
            self.table.item(i, 0).data(QtCore.Qt.ItemDataRole.UserRole) for i in indices
        }
        latest = self.api("/stock")
        stocks = [r for r in latest if r["id"] in ids]
        if len(stocks) != len(ids):
            raise ValueError("Обновите таблицу: часть позиций не найдена")
        if any(qtotal([r["unassigned"], r["returned"]]) <= 0 for r in stocks):
            raise ValueError("Снимите выделение с позиций без доступного остатка")
        from .stock_dialog import StockMoveDialog
        from .batch_dialog import BatchStockDialog

        self._stock_dialog_active = True
        try:
            dialog = (
                StockMoveDialog(self, stocks[0])
                if len(stocks) == 1
                else BatchStockDialog(self, stocks)
            )
            if dialog.exec() == W.QDialog.DialogCode.Accepted:
                self.refresh()
        finally:
            self._stock_dialog_active = False

    def import_file(self):
        path, _ = W.QFileDialog.getOpenFileName(
            self, "Выгрузка 1С", "", "Excel (*.xlsx);;CSV (*.csv)"
        )
        if not path:
            return
        from .import_dialog import ImportDialog

        ImportDialog(self, path).exec()

    def add_location(self):
        self.open_directories()

    def allocate(self):
        self.view_tabs.setCurrentIndex(0)
        self.refresh()

    def move(self, back=False):
        if self.view_tabs.currentIndex() == 0:
            if back:
                raise ValueError(
                    "Для возврата выберите карточки на вкладке «Имущество по складам»"
                )
            self.open_stock_action()
            return
        rows = self.selected()
        if not rows:
            raise ValueError("Выделите карточки")
        if len({r["location"] for r in rows}) != 1:
            raise ValueError("Выберите карточки из одного места")
        from .movement_dialog import MoveAssetsDialog

        if MoveAssetsDialog(self, rows, back).exec() == W.QDialog.DialogCode.Accepted:
            self.refresh()

    def card(self):
        rows = self.selected()
        if not rows:
            return
        r = self.api("/assets/" + rows[0]["id"])
        history = self.api("/history/" + r["id"])
        from .asset_dialog import AssetDialog

        AssetDialog(self, r, history).exec()

    def inventory(self):
        from .inventory_dialog import SessionsDialog

        SessionsDialog(self).exec()

    def session_window(self):
        from .inventory_dialog import InventoryDialog

        if self.sid:
            InventoryDialog(self, self.sid, "Текущая сессия").exec()

    def scan_code(self):
        try:
            asset = self.api("/resolve", {"text": self.scan.text().strip()})
            if self.sid:
                checks = self.api("/sessions/" + self.sid)
                row = next((r for r in checks if r["asset_id"] == asset["id"]), None)
                if not row:
                    raise ValueError(
                        "Имущество вне состава сессии. Автоматическая отметка запрещена"
                    )
                actual = self.scan_location
                if not self.quick.isChecked():
                    if (
                        W.QMessageBox.question(
                            self,
                            "Проверка",
                            asset["number"]
                            + " · подтвердить наличие в выбранном фактическом месте?",
                        )
                        != W.QMessageBox.StandardButton.Yes
                    ):
                        return
                self.command(
                    "/sessions/" + self.sid + "/check",
                    {
                        "asset_id": asset["id"],
                        "version": row["version"],
                        "status": "В наличии"
                        if actual == json.loads(row["snapshot"])["location"]
                        else "Обнаружено в другом месте",
                        "actual": actual,
                        "fast": self.quick.isChecked(),
                    },
                )
                self.statusBar().showMessage("В наличии: " + asset["number"])
            else:
                from .asset_dialog import AssetDialog

                AssetDialog(
                    self,
                    self.api("/assets/" + asset["id"]),
                    self.api("/history/" + asset["id"]),
                ).exec()
        finally:
            self.scan.clear()
            self.scan.setFocus()

    def labels(self):
        rows = self.selected()
        if not rows:
            raise ValueError("Выделите карточки для этикеток")
        template, ok = W.QInputDialog.getItem(
            self, "Этикетки", "Размер", ["58x40", "40x30", "A4"], editable=False
        )
        if not ok:
            return
        path, _ = W.QFileDialog.getSaveFileName(
            self, "Сохранить PDF", "labels.pdf", "PDF (*.pdf)"
        )
        if path:
            Path(path).write_bytes(
                self.api(
                    "/labels", {"assets": [r["id"] for r in rows], "template": template}
                )
            )

    def displayed_rows(self):
        by_id = {row["id"]: row for row in self.visible}
        return [
            by_id[self.table.item(i, 0).data(QtCore.Qt.ItemDataRole.UserRole)]
            for i in range(self.table.rowCount())
        ]

    def clear_filters(self):
        self.column_filters[self.view_tabs.currentIndex()].clear()
        self.search.clear()
        self.populate()

    def remove_filter(self, key):
        self.column_filters[self.view_tabs.currentIndex()].pop(key, None)
        self.populate()

    def filter_presets(self):
        key = "presets/" + str(self.view_tabs.currentIndex())
        presets = json.loads(self.preferences.value(key, "{}"))
        menu = W.QMenu(self)
        save = menu.addAction("Сохранить текущий набор…")
        menu.addSeparator()
        choices = {menu.addAction(name): name for name in sorted(presets)}
        clear = menu.addAction("Удалить все сохранённые наборы…") if presets else None
        action = menu.exec(self.mapToGlobal(self.rect().center()))
        if action is None:
            return
        if action == save:
            name, ok = W.QInputDialog.getText(
                self, "Сохранить фильтр", "Название набора:"
            )
            if not ok or not name.strip():
                return
            presets[name.strip()] = {
                "query": self.search.text(),
                "filters": {
                    k: sorted(v)
                    for k, v in self.column_filters[
                        self.view_tabs.currentIndex()
                    ].items()
                },
            }
            self.preferences.setValue(key, json.dumps(presets, ensure_ascii=False))
        elif action == clear:
            if (
                W.QMessageBox.question(
                    self,
                    "Удаление наборов",
                    "Удалить сохранённые наборы для текущего раздела?",
                )
                == W.QMessageBox.StandardButton.Yes
            ):
                self.preferences.remove(key)
        elif action in choices:
            preset = presets[choices[action]]
            self.column_filters[self.view_tabs.currentIndex()] = {
                k: set(v) for k, v in preset["filters"].items()
            }
            self.search.setText(preset["query"])
            self.populate()

    def save_table_state(self):
        if self._header_view is not None:
            self.preferences.setValue(
                "table/" + ("stock" if self._header_view else "assets"),
                self.table.horizontalHeader().saveState(),
            )

    def configure_columns(self):
        dialog = W.QDialog(self)
        dialog.setWindowTitle("Отображаемые колонки")
        layout = W.QVBoxLayout(dialog)
        checks = []
        for i in range(self.table.columnCount()):
            cb = W.QCheckBox(
                self.table.horizontalHeaderItem(i).text().replace("\n", " ")
            )
            cb.setChecked(not self.table.isColumnHidden(i))
            layout.addWidget(cb)
            checks.append(cb)
        buttons = W.QDialogButtonBox(
            W.QDialogButtonBox.StandardButton.Ok
            | W.QDialogButtonBox.StandardButton.Cancel
        )
        layout.addWidget(buttons)
        buttons.accepted.connect(dialog.accept)
        buttons.rejected.connect(dialog.reject)
        if dialog.exec() == W.QDialog.DialogCode.Accepted:
            if not any(cb.isChecked() for cb in checks):
                raise ValueError("Оставьте хотя бы одну колонку")
            for i, cb in enumerate(checks):
                self.table.setColumnHidden(i, not cb.isChecked())
            self.save_table_state()

    def update_selection(self):
        if not self.current_user or not hasattr(self, "table"):
            return
        count = len(self.table.selectionModel().selectedRows())
        zero = self.view_tabs.currentIndex() == 0
        permissions = self.current_user.get("permissions", [])
        admin = self.current_user.get("is_admin")
        self.buttons["Переместить"].setEnabled(
            bool(
                count
                and (
                    admin
                    or (
                        "allocate" in permissions or "move" in permissions
                        if zero
                        else "move" in permissions
                    )
                )
            )
        )
        self.buttons["Вернуть"].setEnabled(
            bool(count and not zero and (admin or "return" in permissions))
        )
        self.buttons["Этикетки PDF"].setEnabled(
            bool(count and not zero and (admin or "labels" in permissions))
        )
        self.buttons["Переместить"].setToolTip("Выбрано строк: " + str(count))

    def open_documents(self):
        from .documents_dialog import DocumentsDialog

        DocumentsDialog(self).exec()

    def create_backup(self):
        from .backup_dialog import BackupDialog

        BackupDialog(self).exec()

    def refresh_settings(self):
        value, ok = W.QInputDialog.getInt(
            self,
            "Автоматическое обновление",
            "Интервал в секундах (0 — отключить):",
            int(self.preferences.value("refresh_seconds", 15)),
            0,
            86400,
        )
        if ok:
            self.preferences.setValue("refresh_seconds", value)
            if hasattr(self, "refresh_timer"):
                self.refresh_timer.stop()
                if value:
                    self.refresh_timer.start(value * 1000)

    def show_audit(self):
        rows = self.api("/audit")
        dialog = W.QDialog(self)
        dialog.setWindowTitle("Журнал действий")
        dialog.resize(950, 600)
        layout = W.QVBoxLayout(dialog)
        table = W.QTableWidget(len(rows), 4)
        table.setHorizontalHeaderLabels(
            ["Время", "Пользователь", "Действие", "Сведения"]
        )
        table.setEditTriggers(W.QAbstractItemView.EditTrigger.NoEditTriggers)
        for i, row in enumerate(rows):
            for j, key in enumerate(("at", "user_id", "action", "data")):
                table.setItem(i, j, W.QTableWidgetItem(str(row[key])))
        table.resizeColumnsToContents()
        layout.addWidget(table)
        close = W.QPushButton("Закрыть")
        close.clicked.connect(dialog.accept)
        layout.addWidget(close)
        dialog.exec()

    @staticmethod
    def export_value(key, value):
        if key in (
            "quantity",
            "initial",
            "unassigned",
            "returned",
            "allocated",
            "available",
        ):
            q = quantity(value)
            if q == q.to_integral_value():
                return int(q) if len(q.as_tuple().digits) <= 15 else str(value)
            return float(q) if len(q.as_tuple().digits) <= 15 else str(value)
        return value

    def export(self):
        from .security import require

        self.current_user = self.api("/me")
        require(self.current_user, "export")
        from openpyxl import Workbook

        path, _ = W.QFileDialog.getSaveFileName(
            self, "Экспорт текущего поиска", "inventory.xlsx", "Excel (*.xlsx)"
        )
        if not path:
            return
        definitions = STOCK_COLUMNS if self.view_tabs.currentIndex() == 0 else COLUMNS
        header = self.table.horizontalHeader()
        columns = [
            definitions[header.logicalIndex(i)]
            for i in range(header.count())
            if not self.table.isColumnHidden(header.logicalIndex(i))
        ]
        wb = Workbook()
        ws = wb.active
        ws.title = (
            "Нулевой склад" if self.view_tabs.currentIndex() == 0 else "Имущество"
        )
        ws.append([x[1] for x in columns])
        for r in self.displayed_rows():
            ws.append(
                [
                    self.export_value(k, r.get(k, 1 if k == "quantity" else ""))
                    for k, _ in columns
                ]
            )
            for cell in ws[ws.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        from openpyxl.styles import Alignment

        for cell in ws[1]:
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[1].height = 32
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for col in range(1, len(columns) + 1):
            ws.column_dimensions[
                __import__("openpyxl").utils.get_column_letter(col)
            ].width = 25
        meta = wb.create_sheet("Параметры")
        meta.append(
            [
                "Режим",
                "Нулевой склад" if self.view_tabs.currentIndex() == 0 else "Поштучный",
            ]
        )
        meta.append(["Поиск", self.search.text()])
        meta["B2"].data_type = "s"
        meta.append(["Строк", len(self.visible)])
        meta.append(
            [
                "Фильтры",
                json.dumps(
                    {
                        k: sorted(v)
                        for k, v in self.column_filters[
                            self.view_tabs.currentIndex()
                        ].items()
                    },
                    ensure_ascii=False,
                ),
            ]
        )
        meta["B4"].data_type = "s"
        wb.save(path)


def main():
    app = W.QApplication(sys.argv)
    window = Client()
    window.show()
    QtCore.QTimer.singleShot(0, lambda: window.safe(window.login))
    sys.exit(app.exec())


if __name__ == "__main__":
    main()
