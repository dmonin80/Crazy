"""Dedicated inventory workspace with scanning, progress, discrepancies and export."""

import json
from PySide6 import QtWidgets as W
from .reference_widgets import ReferenceCombo


class SessionsDialog(W.QDialog):
    def __init__(self, client):
        super().__init__(client)
        self.client = client
        self.setWindowTitle("Инвентаризации")
        self.resize(850, 560)
        layout = W.QVBoxLayout(self)
        title = W.QLabel("Инвентаризации имущества")
        title.setStyleSheet("font-size:20px;font-weight:600;color:#17436b")
        layout.addWidget(title)
        self.sessions = client.api("/sessions")
        self.list = W.QListWidget()
        layout.addWidget(self.list)
        for row in self.sessions:
            self.list.addItem(
                row["name"] + " · " + row["state"] + " · " + row["at"][:10]
            )
        form = W.QFormLayout()
        layout.addLayout(form)
        self.name = W.QLineEdit()
        self.name.setPlaceholderText("Например: Офис — октябрь 2026")
        form.addRow("Новая сессия", self.name)
        self.locations = client.api("/locations")
        self.location = ReferenceCombo(self.locations, "Все склады")
        form.addRow("Состав сессии", self.location)
        bar = W.QHBoxLayout()
        layout.addLayout(bar)
        create = W.QPushButton("Создать")
        create.clicked.connect(lambda: client.safe(self.create))
        bar.addWidget(create)
        open_button = W.QPushButton("Открыть выбранную")
        open_button.clicked.connect(lambda: client.safe(self.open_selected))
        bar.addWidget(open_button)
        self.list.itemDoubleClicked.connect(lambda _: client.safe(self.open_selected))
        close = W.QPushButton("Закрыть")
        close.clicked.connect(self.reject)
        bar.addWidget(close)

    def create(self):
        if not self.name.text().strip():
            raise ValueError("Укажите название инвентаризации")
        result = self.client.command(
            "/sessions",
            {"name": self.name.text().strip(), "location": self.location.selected_id()},
        )
        self.open_session(result["id"], self.name.text().strip(), "В работе")

    def open_selected(self):
        if self.list.currentRow() < 0:
            raise ValueError("Выберите инвентаризацию")
        row = self.sessions[self.list.currentRow()]
        self.open_session(row["id"], row["name"], row["state"])

    def open_session(self, sid, name, state):
        InventoryDialog(self.client, sid, name, state).exec()
        self.accept()


class InventoryDialog(W.QDialog):
    def __init__(self, client, sid, name, state="В работе"):
        super().__init__(client)
        self.client = client
        self.sid = sid
        self.state = state
        self.rows = []
        self.setWindowTitle("Инвентаризация · " + name)
        self.resize(1150, 760)
        layout = W.QVBoxLayout(self)
        heading = W.QLabel(name)
        heading.setStyleSheet("font-size:20px;font-weight:600;color:#17436b")
        layout.addWidget(heading)
        self.progress = W.QProgressBar()
        layout.addWidget(self.progress)
        scanbar = W.QHBoxLayout()
        layout.addLayout(scanbar)
        self.scan = W.QLineEdit()
        self.scan.setPlaceholderText(
            "Сканируйте QR-код или введите инвентарный номер → Enter"
        )
        scanbar.addWidget(self.scan, 2)
        self.locations = client.api("/locations")
        names = {r["id"]: r["name"] for r in self.locations}
        self.by_id = {
            r["id"]: (names.get(r["parent"], "") + " / " if r.get("parent") else "")
            + r["name"]
            for r in self.locations
        }
        choices = [dict(r, name=self.by_id[r["id"]]) for r in self.locations]
        self.actual = ReferenceCombo(choices, "Выберите фактическое место")
        scanbar.addWidget(self.actual, 1)
        self.quick = W.QCheckBox("Быстрая отметка")
        scanbar.addWidget(self.quick)
        self.scan.returnPressed.connect(lambda: client.safe(self.scan_asset))
        self.feedback = W.QLabel("Выберите фактическое место и начните сканирование.")
        self.feedback.setWordWrap(True)
        self.feedback.setStyleSheet("padding:12px;font-size:16px;background:#eef4fa")
        layout.addWidget(self.feedback)
        filterbar = W.QHBoxLayout()
        layout.addLayout(filterbar)
        self.search = W.QLineEdit()
        self.search.setPlaceholderText("Поиск по номеру и наименованию")
        filterbar.addWidget(self.search, 1)
        self.filter = W.QComboBox()
        self.filter.addItems(
            [
                "Все",
                "Не проверено",
                "В наличии",
                "Не найдено",
                "Обнаружено в другом месте",
            ]
        )
        filterbar.addWidget(self.filter)
        self.search.textChanged.connect(self.populate)
        self.filter.currentTextChanged.connect(self.populate)
        self.table = W.QTableWidget(0, 5)
        self.table.setHorizontalHeaderLabels(
            ["Номер", "Наименование", "Учётное место", "Результат", "Фактическое место"]
        )
        self.table.setSelectionBehavior(
            W.QAbstractItemView.SelectionBehavior.SelectRows
        )
        self.table.setEditTriggers(W.QAbstractItemView.EditTrigger.NoEditTriggers)
        self.table.setAlternatingRowColors(True)
        layout.addWidget(self.table, 1)
        bar = W.QHBoxLayout()
        layout.addLayout(bar)
        self.edit_buttons = []
        for caption, fn, edit in [
            ("Отметить вручную", self.mark, True),
            ("Находка", self.finding, True),
            ("Обновить", self.refresh, False),
            ("Экспорт ведомости", self.export, False),
            ("Завершить сессию", self.finish, True),
        ]:
            button = W.QPushButton(caption)
            button.clicked.connect(lambda checked=False, f=fn: client.safe(f))
            bar.addWidget(button)
            if edit:
                self.edit_buttons.append(button)
        close = W.QPushButton("Закрыть")
        close.clicked.connect(self.accept)
        bar.addWidget(close)
        self.refresh()

    def refresh(self):
        self.rows = self.client.api("/sessions/" + self.sid)
        session = next(
            (s for s in self.client.api("/sessions") if s["id"] == self.sid), None
        )
        if session:
            self.state = session["state"]
        closed = self.state != "В работе"
        for b in self.edit_buttons:
            b.setEnabled(not closed)
        self.scan.setEnabled(not closed)
        self.quick.setEnabled(not closed)
        if closed:
            self.feedback.setText(
                "Сессия завершена. Доступны просмотр и экспорт ведомости."
            )
        self.populate()

    def populate(self, *_):
        query = self.search.text().strip().casefold()
        status = self.filter.currentText()
        self.visible = [
            r
            for r in self.rows
            if (status == "Все" or r["status"] == status)
            and query
            in (
                json.loads(r["snapshot"])["number"]
                + " "
                + json.loads(r["snapshot"])["name"]
            ).casefold()
        ]
        self.table.setRowCount(len(self.visible))
        for i, r in enumerate(self.visible):
            snap = json.loads(r["snapshot"])
            for j, value in enumerate(
                [
                    snap["number"],
                    snap["name"],
                    snap["location_name"],
                    r["status"],
                    self.by_id.get(r["actual"], ""),
                ]
            ):
                self.table.setItem(i, j, W.QTableWidgetItem(value))
        self.table.resizeColumnsToContents()
        self.table.horizontalHeader().setStretchLastSection(True)
        checked = sum(r["status"] != "Не проверено" for r in self.rows)
        issues = sum(
            r["status"] in ("Не найдено", "Обнаружено в другом месте")
            for r in self.rows
        )
        self.progress.setRange(0, max(1, len(self.rows)))
        self.progress.setValue(checked)
        self.progress.setFormat(
            f"Проверено {checked} из {len(self.rows)} · Расхождений: {issues}"
        )

    def scan_asset(self):
        text = self.scan.text().strip()
        if not text:
            return
        try:
            actual = self.actual.selected_id()
            if actual is None:
                raise ValueError("Выберите фактическое место")
            if (
                not self.quick.isChecked()
                and W.QMessageBox.question(
                    self, "Подтверждение", "Подтвердить наличие: " + text + "?"
                )
                != W.QMessageBox.StandardButton.Yes
            ):
                return
            result = self.client.command(
                "/sessions/" + self.sid + "/scan",
                {"text": text, "actual": actual, "fast": self.quick.isChecked()},
            )
            self.refresh()
            self.feedback.setText(
                "✓ "
                + result["number"]
                + " · "
                + result["status"]
                + (
                    " · Внимание: карточка перемещалась после начала сессии"
                    if result.get("movement_warning")
                    else ""
                )
            )
        except Exception as exc:
            self.feedback.setText("⚠ " + str(exc))
        finally:
            self.scan.clear()
            self.scan.setFocus()

    def mark(self):
        index = self.table.currentRow()
        if index < 0:
            raise ValueError("Выберите строку")
        row = self.visible[index]
        status, ok = W.QInputDialog.getItem(
            self,
            "Результат",
            "Статус",
            ["В наличии", "Не найдено", "Обнаружено в другом месте", "Не проверено"],
            editable=False,
        )
        if not ok:
            return
        actual = (
            self.actual.selected_id()
            if status in ("В наличии", "Обнаружено в другом месте")
            else None
        )
        if status in ("В наличии", "Обнаружено в другом месте") and not actual:
            raise ValueError("Выберите фактическое место вверху окна")
        comment, ok = W.QInputDialog.getText(self, "Результат", "Комментарий:")
        if not ok:
            return
        result = self.client.command(
            "/sessions/" + self.sid + "/check",
            {
                "asset_id": row["asset_id"],
                "version": row["version"],
                "status": status,
                "actual": actual,
                "comment": comment,
                "fast": False,
            },
        )
        self.refresh()
        if result.get("movement_warning"):
            self.feedback.setText(
                "⚠ Результат сохранён. Карточка перемещалась после начала сессии: проверьте историю."
            )

    def finding(self):
        description, ok = W.QInputDialog.getMultiLineText(
            self, "Находка", "Номер, описание и фактическое место:"
        )
        if ok and description.strip():
            self.client.command(
                "/sessions/" + self.sid + "/finding",
                {"description": description.strip()},
            )
            self.feedback.setText("✓ Находка зарегистрирована")

    def finish(self):
        self.refresh()
        if self.state != "В работе":
            return
        remaining = sum(r["status"] == "Не проверено" for r in self.rows)
        issues = sum(
            r["status"] in ("Не найдено", "Обнаружено в другом месте")
            for r in self.rows
        )
        if (
            W.QMessageBox.question(
                self,
                "Завершение",
                f"Не проверено: {remaining}. Расхождений: {issues}. Завершить сессию?",
            )
            != W.QMessageBox.StandardButton.Yes
        ):
            return
        self.client.command("/sessions/" + self.sid + "/finish", {"confirm": True})
        self.state = "Завершена"
        if self.client.sid == self.sid:
            self.client.sid = None
            self.client.session_label.setText("Сессия завершена")
            self.client.quick.setChecked(False)
        self.refresh()

    def export(self):
        from openpyxl import Workbook

        path, _ = W.QFileDialog.getSaveFileName(
            self, "Ведомость", "inventory-check.xlsx", "Excel (*.xlsx)"
        )
        if not path:
            return
        rows = self.client.api("/sessions/" + self.sid)
        wb = Workbook()
        ws = wb.active
        ws.title = "Ведомость"
        ws.append(
            [
                "Номер",
                "Наименование",
                "Учётное место",
                "Результат",
                "Фактическое место",
                "Комментарий",
            ]
        )
        for r in rows:
            snap = json.loads(r["snapshot"])
            ws.append(
                [
                    snap["number"],
                    snap["name"],
                    snap["location_name"],
                    r["status"],
                    self.by_id.get(r["actual"], ""),
                    r["comment"],
                ]
            )
            for cell in ws[ws.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for col in "ABCDEF":
            ws.column_dimensions[col].width = 28
        wb.save(path)
