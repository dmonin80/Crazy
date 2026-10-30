"""Read 1C exports without changing original item names or codes."""

import csv
import hashlib
import re
from datetime import date, datetime
from decimal import InvalidOperation
from io import BytesIO, StringIO
from pathlib import Path
from .limits import MAX_FILE_BYTES, MAX_EXPANDED_BYTES, MAX_ROWS, MAX_COLUMNS
from .quantities import (
    quantity as parse_quantity,
    number as qnumber,
    total as qtotal,
    mode as qmode,
)

FIELDS = {
    "name": "Номенклатура",
    "code": "Код",
    "unit": "Единица измерения",
    "quantity": "Количество",
}


def normalized(value):
    return re.sub(r"[^a-zа-я0-9]", "", str(value or "").casefold().replace("ё", "е"))


ALIASES = {
    "name": {"номенклатура", "наименование", "наименованиеноменклатуры"},
    "code": {"код", "код1с", "кодноменклатуры", "кодноменклатуры1с"},
    "unit": {"единицаизмерения", "единицаизмерений", "едизм", "едизмерения", "ед"},
    "quantity": {"количество", "колво", "общееколичество"},
}


def column_letter(index):
    from openpyxl.utils import get_column_letter

    return get_column_letter(index + 1)


def display(value):
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value)


class Table:
    def __init__(self, rows, sheets, sheet, formulas=None, numeric_codes=None):
        self.rows = rows
        self.sheets = sheets
        self.sheet = sheet
        self.formulas = formulas or set()
        self.numeric_codes = numeric_codes or set()
        self.width = max((len(values) for _, values in rows), default=0)


def read_table(data, filename, sheet=None, encoding="utf-8-sig", delimiter=";"):
    if len(data) > MAX_FILE_BYTES:
        raise ValueError("Файл больше 20 МБ. Разделите выгрузку на части.")
    extension = Path(filename).suffix.casefold()
    if extension == ".xlsx":
        from openpyxl import load_workbook
        from zipfile import BadZipFile, ZipFile

        try:
            with ZipFile(BytesIO(data)) as archive:
                if (
                    len(archive.infolist()) > 2000
                    or sum(i.file_size for i in archive.infolist()) > MAX_EXPANDED_BYTES
                ):
                    raise ValueError(
                        "Слишком большой распакованный XLSX (максимум 100 МБ и 2000 частей)"
                    )
            wb = load_workbook(BytesIO(data), data_only=True, read_only=True)
            fw = load_workbook(BytesIO(data), data_only=False, read_only=True)
        except (BadZipFile, OSError) as exc:
            raise ValueError("Файл XLSX повреждён или имеет другой формат") from exc
        try:
            if sheet and sheet not in wb.sheetnames:
                raise ValueError("Лист не найден: " + sheet)
            ws = wb[sheet] if sheet else wb.active
            formulas, numeric = set(), set()
            raw = []
            if (ws.max_row or 0) > MAX_ROWS or (ws.max_column or 0) > MAX_COLUMNS:
                raise ValueError(
                    "Допустимо до 100 000 строк и 256 колонок. Удалите лишнее форматирование листа."
                )
            formula_rows = fw[ws.title].iter_rows()
            for row_number, cells in enumerate(ws.iter_rows(), 1):
                formula_cells = next(formula_rows)
                if not cells:
                    continue
                values = []
                for column_number, cell in enumerate(cells, 1):
                    value = cell.value
                    position = (row_number, column_number - 1)
                    if (
                        value is None
                        and formula_cells[column_number - 1].data_type == "f"
                    ):
                        formulas.add(position)
                    if isinstance(value, (int, float)) and not isinstance(value, bool):
                        numeric.add(position)
                        # Only a simple Excel zero-mask unambiguously restores leading zeroes.
                        if (
                            cell.number_format
                            and set(cell.number_format) == {"0"}
                            and float(value).is_integer()
                        ):
                            value = format(
                                int(value), "0" + str(len(cell.number_format)) + "d"
                            )
                    values.append(value)
                raw.append((row_number, values))
            return Table(raw, wb.sheetnames, ws.title, formulas, numeric)
        finally:
            wb.close()
            fw.close()
    if extension == ".csv":
        if len(delimiter) != 1:
            raise ValueError("Разделитель должен быть одним символом")
        try:
            rows = []
            for number, values in enumerate(
                csv.reader(StringIO(data.decode(encoding)), delimiter=delimiter), 1
            ):
                if number > MAX_ROWS or len(values) > MAX_COLUMNS:
                    raise ValueError("Допустимо до 100 000 строк и 256 колонок")
                rows.append((number, values))
        except (UnicodeError, LookupError) as exc:
            raise ValueError(
                "Не удалось прочитать CSV. Выберите правильную кодировку."
            ) from exc
        return Table(rows, ["CSV"], "CSV")
    raise ValueError("Поддерживаются XLSX и CSV. Файл XLS нужно пересохранить в XLSX.")


def guess_mapping(values):
    mapping = {}
    for field, aliases in ALIASES.items():
        matches = [
            index for index, value in enumerate(values) if normalized(value) in aliases
        ]
        if len(matches) == 1:
            mapping[field] = matches[0]
    return mapping


def guess_header(table):
    best_header, best_mapping = 1, {}
    for row, values in table.rows[:100]:
        mapping = guess_mapping(values)
        if len(mapping) > len(best_mapping):
            best_header, best_mapping = row, mapping
    return best_header, best_mapping


def inspect_source(
    data, filename, sheet=None, header=None, encoding="utf-8-sig", delimiter=";"
):
    table = read_table(data, filename, sheet, encoding, delimiter)
    if header is None:
        header, mapping = guess_header(table)
        # Prefer a full 1C export header over an unrelated active sheet.
        if sheet is None and len(mapping) < 4 and len(table.sheets) > 1:
            for name in table.sheets:
                if name == table.sheet:
                    continue
                other = read_table(data, filename, name, encoding, delimiter)
                other_header, other_mapping = guess_header(other)
                if len(other_mapping) > len(mapping):
                    table, header, mapping = other, other_header, other_mapping
                if len(mapping) == 4:
                    break
    else:
        if type(header) is not int or not 1 <= header <= len(table.rows):
            raise ValueError("Строка заголовков находится вне листа")
        mapping = guess_mapping(table.rows[header - 1][1])
    header_values = table.rows[header - 1][1] if table.rows else []
    columns = [
        {
            "index": i,
            "letter": column_letter(i),
            "heading": display(header_values[i]) if i < len(header_values) else "",
        }
        for i in range(table.width)
    ]
    return {
        "sheet": table.sheet,
        "sheets": table.sheets,
        "header": header,
        "mapping": mapping,
        "columns": columns,
        "row_count": len(table.rows),
        "sample_rows": [
            {"row": n, "values": [display(v) for v in values]}
            for n, values in table.rows[header : header + 30]
        ],
    }


def preview(
    data,
    filename,
    mapping,
    header=1,
    sheet=None,
    encoding="utf-8-sig",
    delimiter=";",
    merge=False,
    excluded_rows=None,
):
    table = read_table(data, filename, sheet, encoding, delimiter)
    if type(header) is not int or not 1 <= header <= len(table.rows):
        raise ValueError("Укажите существующую строку заголовков")
    if (
        not isinstance(mapping, dict)
        or set(mapping) != set(FIELDS)
        or any(
            type(i) is not int or i < 0 or i >= table.width for i in mapping.values()
        )
        or len(set(mapping.values())) != 4
    ):
        raise ValueError(
            "Выберите четыре разные существующие колонки: номенклатура, код, единица измерения, количество"
        )
    excluded_rows = [] if excluded_rows is None else excluded_rows
    if (
        not isinstance(excluded_rows, list)
        or any(
            type(n) is not int or not header < n <= len(table.rows)
            for n in excluded_rows
        )
        or len(set(excluded_rows)) != len(excluded_rows)
    ):
        raise ValueError(
            "Список исключений должен содержать номера существующих строк данных без повторов"
        )
    excluded_set = set(excluded_rows)
    records, errors, warnings, duplicates, skipped, excluded = {}, [], [], [], [], []
    warned_codes = set()

    def issue(line, field, reason, values):
        column = mapping[field] if field else None
        return {
            "row": line,
            "field": FIELDS.get(field, ""),
            "cell": column_letter(column) + str(line) if column is not None else "",
            "value": display(values[column])
            if column is not None and column < len(values)
            else "",
            "reason": reason,
        }

    for line, values in table.rows[header:]:
        if line in excluded_set:
            excluded.append(
                {
                    "row": line,
                    "values": {
                        k: display(values[i]) if i < len(values) else ""
                        for k, i in mapping.items()
                    },
                }
            )
            continue
        selected = {
            key: values[index] if index < len(values) else None
            for key, index in mapping.items()
        }
        if not any(
            value is not None and str(value).strip() for value in selected.values()
        ):
            if any(value is not None and str(value).strip() for value in values):
                skipped.append(
                    {
                        "row": line,
                        "reason": "Служебная строка без данных в четырёх выбранных колонках",
                    }
                )
            continue
        header_fields = guess_mapping(values)
        if all(header_fields.get(key) == index for key, index in mapping.items()):
            skipped.append({"row": line, "reason": "Повторная строка заголовков"})
            continue
        label = normalized(selected["name"])
        if (label.startswith("итого") or label.startswith("всего")) and not display(
            selected["code"]
        ).strip():
            skipped.append(
                {"row": line, "reason": "Строка итогов без кода номенклатуры"}
            )
            continue
        row_errors = []
        for field, value in selected.items():
            if (line, mapping[field]) in table.formulas:
                row_errors.append(
                    issue(
                        line,
                        field,
                        "Формула без сохранённого результата. Пересохраните файл в Excel с рассчитанными значениями.",
                        values,
                    )
                )
            elif value is None or not str(value).strip():
                row_errors.append(
                    issue(
                        line,
                        field,
                        "Не заполнено обязательное поле «" + FIELDS[field] + "»",
                        values,
                    )
                )
        if row_errors:
            errors.extend(row_errors)
            continue
        try:
            text = (
                str(selected["quantity"])
                .strip()
                .replace("\u00a0", "")
                .replace("\u202f", "")
                .replace(" ", "")
                .replace(",", ".")
            )
            quantity = parse_quantity(text)
        except (ValueError, InvalidOperation, OverflowError) as exc:
            errors.append(
                issue(
                    line,
                    "quantity",
                    str(exc)
                    if isinstance(exc, ValueError)
                    else "Некорректное количество",
                    values,
                )
            )
            continue
        record = {key: str(selected[key]) for key in ("name", "code", "unit")}
        record.update(
            quantity=qnumber(quantity),
            tracking_mode=qmode(record["unit"], quantity),
            rows=[line],
        )
        code = record["code"]
        if (line, mapping["code"]) in table.numeric_codes and code not in warned_codes:
            warnings.append(
                issue(
                    line,
                    "code",
                    "Код хранится числом в Excel. Проверьте ведущие нули и точность; уже утраченные символы восстановить невозможно.",
                    values,
                )
            )
            warned_codes.add(code)
        if code in records:
            old = records[code]
            if (old["name"], old["unit"]) != (record["name"], record["unit"]):
                errors.append(
                    issue(
                        line,
                        "code",
                        "Одинаковый код с другим наименованием или единицей измерения; исходная строка "
                        + str(old["rows"][0]),
                        values,
                    )
                )
                continue
            duplicates.append({"code": code, "row": line, "first_row": old["rows"][0]})
            old["quantity"] = qnumber(qtotal([old["quantity"], record["quantity"]]))
            if record["tracking_mode"] == "quantity":
                old["tracking_mode"] = "quantity"
            old["rows"].append(line)
        else:
            records[code] = record
    if duplicates and not merge:
        errors.append(
            {
                "row": None,
                "field": "Код",
                "cell": "",
                "value": "",
                "reason": "Найдены одинаковые коды с одинаковыми реквизитами. Для объединения включите соответствующий флажок и повторите проверку.",
            }
        )
    if not records:
        errors.append(
            {
                "row": None,
                "field": "",
                "cell": "",
                "value": "",
                "reason": "Нет строк номенклатуры для импорта",
            }
        )
    return {
        "rows": list(records.values()),
        "errors": errors,
        "warnings": warnings,
        "duplicates": duplicates,
        "hash": hashlib.sha256(data).hexdigest(),
        "units": qnumber(qtotal(r["quantity"] for r in records.values())),
        "excluded": excluded,
        "skipped": skipped,
        "sheet": table.sheet,
        "header": header,
        "mapping": mapping,
    }
